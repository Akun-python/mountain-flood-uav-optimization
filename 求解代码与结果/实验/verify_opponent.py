# -*- coding: utf-8 -*-
"""对方方案自动复核工具（读竞赛结果提交模板 xlsx 的 Q2 表做四项核验）。
用法： python verify_opponent.py <对方结果提交模板.xlsx>
原始校验项：
  1) 投送区一致性 —— 逐箱交付的"服务区"必须 = 箱属区
  2) 硬时限        —— 首批截止 / 医疗期望；其余按期望衡量加权迟到
  3) 电池周转      —— 每电池链条 SOC 回收 + 两阶段充电，不重叠、不超额(≤A6/B4/C4)
  4) 完工/架次/能耗口径 —— 汇总申报指标（能耗需 Q1 组批才能逐三重算，此处给口径提示）
"""
import sys, os, json
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '代码'))
from core import Data, charge_time

HERE = os.path.dirname(os.path.abspath(__file__))
OUTD = os.path.join(HERE, '..', '结果', '进化_v25')
data = Data()


def box_area(b):
    return b.split('-')[0]


def _col(ws, header_text, required=True):
    for idx, c in enumerate(ws[1], start=1):
        if c.value and str(c.value).strip() == header_text:
            return idx
    if required:
        raise KeyError('找不到列 %r' % header_text)
    return None


def read_rows(ws):
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r is None or all(v is None or str(v).strip() == '' for v in r):
            continue
        yield r


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else input('对方xlsx: ')
    wb = openpyxl.load_workbook(src, data_only=True)
    R = {'src': src, 'errors': [], 'summary': {}}

    # ---------- Q2_逐箱交付 ----------
    if 'Q2_逐箱交付' in wb.sheetnames:
        ws = wb['Q2_逐箱交付']
        cbid = _col(ws, '货箱编号'); cfid = _col(ws, '架次编号')
        carea = _col(ws, '服务区编号'); ct = _col(ws, '交付完成时刻（s）')
        zone_bad, dead_bad = [], []
        late_w = 0.0; seen = set(); nrow = 0
        for r in read_rows(ws):
            nrow += 1
            bid = str(r[cbid - 1]).strip(); sid = str(r[carea - 1]).strip().upper()
            t = float(r[ct - 1]); seen.add(bid)
            if box_area(bid) != sid:
                zone_bad.append((bid, sid))
            bx = data.boxes.get(bid)
            if bx is None:
                R['errors'].append('未知货箱 %s' % bid); continue
            if bx['first_batch'] and t > bx['deadline_first'] + 1e-6:
                dead_bad.append((bid, 'first', round(t), bx['deadline_first']))
            if bx['type'] == '医疗物资' and t > bx['deadline_exp'] + 1e-6:
                dead_bad.append((bid, 'medical', round(t), bx['deadline_exp']))
            late_w += max(0.0, t - bx['deadline_exp']) * bx['priority'] / 24.0
        expect80 = set(data.boxes.keys())
        missing = expect80 - seen
        R['summary']['Q2_逐箱交付行'] = nrow
        R['summary']['投送区≠箱属区'] = len(zone_bad)
        R['summary']['硬时限违规'] = len(dead_bad)
        R['summary']['加权迟到(期望口径)'] = round(late_w, 2)
        R['summary']['缺失货箱'] = len(missing)
        if zone_bad: R['errors'].append('投送区≠箱属区: ' + str(zone_bad[:12]))
        if dead_bad: R['errors'].append('硬时限违规: ' + str(dead_bad[:12]))
        if missing: R['errors'].append('缺箱: ' + str(list(missing)[:12]))

    # ---------- Q2_运输架次 ----------
    if 'Q2_运输架次' in wb.sheetnames:
        ws = wb['Q2_运输架次']
        cfid=_col(ws,'架次编号'); cuav=_col(ws,'无人机编号'); cmodel=_col(ws,'机型编号')
        cbat=_col(ws,'电池编号'); cstart=_col(ws,'开始时刻（s）')
        croute=_col(ws,'访问服务区顺序'); cret=_col(ws,'返回O01时刻（s）'); ce=_col(ws,'架次能耗（kWh）')
        rows = list(read_rows(ws))
        # 电池周转
        per_bat = {}
        for r in rows:
            b = str(r[cbat - 1]).strip()
            per_bat.setdefault(b, []).append(r)
        cnt_model = {}
        rmat_errors = []
        for b, rl in per_bat.items():
            rl.sort(key=lambda r: float(r[cstart - 1]))
            model = str(rl[0][cmodel - 1]).strip()
            cnt_model[model] = cnt_model.get(model, 0) + 1
            for a, b2 in zip(rl, rl[1:]):
                e = float(a[ce - 1]); retf = float(a[cret - 1])
                soc = 1 - e / data.uav_types[model]['E_use'] if e <= data.uav_types[model]['E_use'] else 1.0
                ch = charge_time(max(0.0, min(1.0, soc)), data.batteries[model]['T_full'])
                nxt = float(b2[cstart - 1])
                if nxt < retf + ch - 1e-6:
                    rmat_errors.append('电池 %s 充电周转不足 f%s(ret%.0f+ch%.0f)->f%s' % (b, a[cfid-1], retf, ch, b2[cfid-1]))
        inv = {'A': 6, 'B': 4, 'C': 4}
        for g, v in cnt_model.items():
            if v > inv.get(g, 0):
                rmat_errors.append('电池 %s 超额 %d/%d' % (g, v, inv[g]))
        R['summary']['电池用量'] = cnt_model
        R['summary']['运输架次'] = len(rows)
        makespan = max((float(r[cret - 1]) for r in rows), default=0)
        en_sum = sum(float(r[ce - 1]) for r in rows)
        R['summary']['申报完工(最晚返场)'] = round(makespan, 1)
        R['summary']['申报运输能耗合计'] = round(en_sum, 2)
        if rmat_errors: R['errors'].extend(rmat_errors)
        # 机队链趟数
        per_uav = {}
        for r in rows:
            u = str(r[cuav - 1]).strip(); per_uav.setdefault(u, 0); per_uav[u] += 1
        R['summary']['各机架次数'] = {u: n for u, n in sorted(per_uav.items())}
        if R['summary']['各机架次数']:
            mx = max(R['summary']['各机架次数'].values()); nn = sum(R['summary']['各机架次数'].values())
            R['summary']['平均每机趟数'] = round(nn / len(per_uav), 2)

    print(json.dumps(R, ensure_ascii=False, indent=1))
    json.dump(R, open(os.path.join(OUTD, 'opponent_review.json'), 'w', encoding='utf-8'),
              ensure_ascii=False, indent=1)
    print('\nsaved -> 结果/进化_v25/opponent_review.json')


if __name__ == '__main__':
    main()