# -*- coding: utf-8 -*-
"""
D题 山区洪涝灾害下无人机运输与通信协同优化 —— 共享物理核心
包含：数据装载、DEM 采样、几何/航段时间与能耗、电池充电、通信链路判定。

统一口径（附录2/附录3）：
- 平面距离：局部等距投影（以 O01 为原点）。
- 巡航海拔 = 航段沿线 DEM 最高地面高程 + 50 m。
- 航段时间 t = h+/v_up + d/v_c + h-/v_down；下降不单独计能耗。
- 水平能耗：E_hor(q) = E_use * d / L_g(q)；L_g(q) = L0 - (L0-LF)*(q/Q)^1.5。
- 爬升附加能耗：E_up(q) = (m0+q)*g*h+ / (eta_climb*3.6e6)  [kWh]。
- 返航安全余量：E_total <= (1-rho)*E_use。
- 充电两阶段模型。
- 通信：LOS 遮挡 + FSPL + 双向门限。
"""
import os
import math
import json
import numpy as np
from scipy.io import loadmat

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..',
                    '第二十三届中国研究生数学建模竞赛 - 中文题目', '中文题目', 'D题', '数据')
DATA_CSV = os.path.join(BASE, '无人机应急物资运输基础数据')
GIS = os.path.join(BASE, '镇龙乡地理空间数据', '镇龙乡及周边地理数据')

G = 9.80665  # m/s^2
J_PER_KWH = 3.6e6

# ----------------------------------------------------------------------------
# 数据装载
# ----------------------------------------------------------------------------
def load_excel(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
    return rows


class Data:
    """装载并解析所有数据文件。"""

    def __init__(self):
        self.load_centers()
        self.load_boxes()
        self.load_uavs()
        self.load_relay()
        self.load_comm()
        self.load_dem()
        self.build_leg_table()

    # ---- 调度中心与服务区 ----
    def load_centers(self):
        rows = load_excel(os.path.join(DATA_CSV, '调度中心与服务区.xlsx'))
        self.centers = {}
        self.areas = {}
        # 第一段: 调度中心
        r = rows[2]  # O01 行
        oid = r[0]
        self.centers[oid] = {'lon': float(r[2]), 'lat': float(r[3]), 'alt': float(r[4])}
        # 第二段: 服务区（空行已过滤，表头为 rows[4]，数据 rows[5:20]）
        for r in rows[5:20]:
            if r[0] is None:
                continue
            sid = r[0]
            self.areas[sid] = {
                'name': r[1], 'lon': float(r[2]), 'lat': float(r[3]),
                'alt': float(r[4]), 'pop': float(r[5]),
            }
        self.OID = oid
        self.lon0 = self.centers[oid]['lon']
        self.lat0 = self.centers[oid]['lat']
        self.lat0_rad = math.radians(self.lat0)

    def xy(self, lon, lat):
        """局部平面坐标 (m)，以 O01 为原点。"""
        dx = (lon - self.lon0) * 111320.0 * math.cos(self.lat0_rad)
        dy = (lat - self.lat0) * 110540.0
        return dx, dy

    def dist_ll(self, lon1, lat1, lon2, lat2):
        x1, y1 = self.xy(lon1, lat1)
        x2, y2 = self.xy(lon2, lat2)
        return math.hypot(x2 - x1, y2 - y1)

    # ---- 物资与货箱 ----
    def load_boxes(self):
        rows = load_excel(os.path.join(DATA_CSV, '物资需求与配送时限.xlsx'))
        self.boxes = {}   # box_id -> dict
        self.area_boxes = {}  # area -> [box_id]
        # 逐箱货箱清单从第二个 sheet 读
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(DATA_CSV, '物资需求与配送时限.xlsx'), data_only=True)
        ws = wb.worksheets[1]
        for r in ws.iter_rows(values_only=True):
            if r[0] is None or not str(r[0]).startswith('S'):
                continue
            bid, sid, mtype = r[0], r[1], r[2]
            box = {
                'area': sid, 'type': mtype,
                'mass': float(r[3]), 'vol': float(r[4]),
                'first_batch': (str(r[5]).strip() == '是'),
                'deadline_first': float(r[6]) if r[6] is not None else math.inf,
                'deadline_exp': float(r[7]) if r[7] is not None else math.inf,
                'priority': float(r[8]),
            }
            self.boxes[bid] = box
            self.area_boxes.setdefault(sid, []).append(bid)
        self.box_ids = list(self.boxes.keys())
        # 服务区聚合需求
        self.area_demand = {}
        for sid, bids in self.area_boxes.items():
            m = sum(self.boxes[b]['mass'] for b in bids)
            v = sum(self.boxes[b]['vol'] for b in bids)
            self.area_demand[sid] = {'mass': m, 'vol': v, 'nbox': len(bids)}

    # ---- 运输无人机 ----
    def load_uavs(self):
        rows = load_excel(os.path.join(DATA_CSV, '运输无人机数据.xlsx'))
        self.uav_types = {}   # g -> params
        for r in rows[2:5]:
            g = r[0]
            self.uav_types[g] = {
                'name': r[1], 'm0': float(r[2]), 'Q': float(r[3]), 'V': float(r[4]),
                'v_c': float(r[5]), 'L0': float(r[6]), 'LF': float(r[7]),
                'E_use': float(r[8]), 'rho': float(r[9]) / 100.0,
                't_prep': float(r[10]), 't_load': float(r[11]),
                't_hand_base': float(r[12]), 't_hand_box': float(r[13]),
                'v_up': float(r[14]), 'v_down': float(r[15]),
                'eta_up': float(r[16]), 'eta_down': float(r[17]),
            }
        # 逐架无人机 (空行已被过滤, rows[7:15] = U01..U08)
        self.uavs = []      # [(uav_id, type)]
        for r in rows[7:15]:
            if r[0] is None:
                continue
            self.uavs.append((r[0], r[1]))
        # 共享电池 (rows[17:20] = A,B,C)
        self.batteries = {}  # g -> {'n': int, 'T_full': s}
        for r in rows[17:20]:
            if r[0] is None:
                continue
            self.batteries[r[0]] = {'n': int(r[1]), 'T_full': float(r[2])}

    # ---- 中继无人机 ----
    def load_relay(self):
        rows = load_excel(os.path.join(DATA_CSV, '中继无人机数据.xlsx'))
        r = rows[2]
        self.relay_type = {
            'name': r[1], 'm0': float(r[2]), 'm_comm': float(r[3]), 'm_takeoff': float(r[4]),
            'v_c': float(r[5]), 'P_cruise': float(r[6]), 'E_use': float(r[7]),
            'rho': float(r[8]) / 100.0, 't_prep': float(r[9]), 't_link': float(r[10]),
            't_turn': float(r[11]), 'v_up': float(r[12]), 'v_down': float(r[13]),
            'eta_up': float(r[14]), 'eta_down': float(r[15]),
            'P_hover': float(r[16]), 'P_comm': float(r[17]), 'h_max': float(r[18]),
        }
        self.relays = [(r[0], r[1]) for r in rows[5:7] if r[0] is not None]
        r0 = rows[9]
        self.relay_energy = {'n': int(r0[1]), 'T_full': float(r0[2])}

    # ---- 通信链路参数 ----
    def load_comm(self):
        rows = load_excel(os.path.join(DATA_CSV, '通信链路参数.xlsx'))
        d = {}
        for r in rows[2:]:
            if r[0] is None:
                continue
            d[(r[0], r[1])] = r[4]
        self.comm = {
            'f_MHz': float(d[('传播参数', '载波频率（MHz）')]),
            'Lsys': float(d[('传播参数', '系统损耗（dB）')]),
            'Lobs': float(d[('传播参数', '地形遮挡附加损耗（dB）')]),
            'Psens': float(d[('接收参数', '接收灵敏度（dBm）')]),
            'M': float(d[('接收参数', '衰落裕量（dB）')]),
            'Pt_uav': float(d[('运输无人机', '发射功率（dBm）')]),
            'G_uav': float(d[('运输无人机', '天线增益（dBi）')]),
            'Pt_acc': float(d[('中继接入端', '发射功率（dBm）')]),
            'G_acc': float(d[('中继接入端', '天线增益（dBi）')]),
            'Pt_bh': float(d[('中继回传端', '发射功率（dBm）')]),
            'G_bh': float(d[('中继回传端', '天线增益（dBi）')]),
            'Pt_gw': float(d[('固定网关 G01', '发射功率（dBm）')]),
            'G_gw': float(d[('固定网关 G01', '天线增益（dBi）')]),
            'h_gw': float(d[('固定网关 G01', '天线离地高度（m）')]),
        }
        c = self.comm
        Pth = c['Psens'] + c['M']
        c['Lmax_uav2gw'] = min(c['Pt_uav'] + c['G_uav'] + c['G_gw'] - c['Lsys'] - Pth,
                               c['Pt_gw'] + c['G_gw'] + c['G_uav'] - c['Lsys'] - Pth)
        c['Lmax_uav2relay'] = min(c['Pt_uav'] + c['G_uav'] + c['G_acc'] - c['Lsys'] - Pth,
                                  c['Pt_acc'] + c['G_acc'] + c['G_uav'] - c['Lsys'] - Pth)
        c['Lmax_relay2gw'] = min(c['Pt_bh'] + c['G_bh'] + c['G_gw'] - c['Lsys'] - Pth,
                                 c['Pt_gw'] + c['G_gw'] + c['G_bh'] - c['Lsys'] - Pth)
        # 网关通信端点
        o = self.centers[self.OID]
        self.gw_pt = {'lon': o['lon'], 'lat': o['lat'], 'z': o['alt'] + c['h_gw']}

    # ---- DEM ----
    def load_dem(self):
        dem_path = os.path.join(GIS, '数字高程模型数据（DEM）', '镇龙乡及周边30米DEM.mat')
        m = loadmat(dem_path)
        self.dem = np.asarray(m['dem'], dtype=float)
        self.dem_nodata = float(np.asarray(m['nodata']).ravel()[0])
        self.dem_lat_max = float(np.asarray(m['latitude']).max())
        self.dem_lon_min = float(np.asarray(m['longitude']).min())
        self.dem_res = 1.0 / 3600.0  # 度/像元
        self.dem[self.dem == self.dem_nodata] = np.nan

    def dem_at(self, lon, lat):
        """双线性插值 DEM 高程（m）；越界或 nodata 返回 nan。"""
        col_f = (lon - self.dem_lon_min) / self.dem_res
        row_f = (self.dem_lat_max - lat) / self.dem_res
        c0 = int(math.floor(col_f)); r0 = int(math.floor(row_f))
        if c0 < 0 or r0 < 0 or c0 + 1 >= self.dem.shape[1] or r0 + 1 >= self.dem.shape[0]:
            return np.nan
        fc, fr = col_f - c0, row_f - r0
        a = self.dem[r0:r0 + 2, c0:c0 + 2]
        if np.isnan(a).any():
            return np.nan
        return (a[0, 0] * (1 - fr) * (1 - fc) + a[0, 1] * (1 - fr) * fc
                + a[1, 0] * fr * (1 - fc) + a[1, 1] * fr * fc)

    def dem_at_many(self, lons, lats):
        """向量化双线性插值 DEM。返回 array（越界/nodata 为 nan）。"""
        lons = np.asarray(lons, dtype=float)
        lats = np.asarray(lats, dtype=float)
        col_f = (lons - self.dem_lon_min) / self.dem_res
        row_f = (self.dem_lat_max - lats) / self.dem_res
        c0 = np.floor(col_f).astype(int)
        r0 = np.floor(row_f).astype(int)
        fc = col_f - c0
        fr = row_f - r0
        H, W = self.dem.shape
        valid = (c0 >= 0) & (r0 >= 0) & (c0 + 1 < W) & (r0 + 1 < H)
        out = np.full(lons.shape, np.nan)
        r0c = np.clip(r0, 0, H - 2)
        c0c = np.clip(c0, 0, W - 2)
        a = self.dem[r0c, c0c]
        b = self.dem[r0c, c0c + 1]
        c = self.dem[r0c + 1, c0c]
        d = self.dem[r0c + 1, c0c + 1]
        top = a * (1 - fc) + b * fc
        bot = c * (1 - fc) + d * fc
        vals = top * (1 - fr) + bot * fr
        vals[(a == self.dem_nodata) | (b == self.dem_nodata) | (c == self.dem_nodata) | (d == self.dem_nodata)] = np.nan
        out[valid] = vals[valid]
        return out

    def corridor_max(self, lon1, lat1, lon2, lat2):
        """航段沿线 DEM 最大高程（含两端点）。"""
        d = self.dist_ll(lon1, lat1, lon2, lat2)
        n = max(3, int(math.ceil(d / 30.0)) + 1)
        vals = []
        for k in range(n + 1):
            t = k / n
            h = self.dem_at(lon1 + (lon2 - lon1) * t, lat1 + (lat2 - lat1) * t)
            if not math.isnan(h):
                vals.append(h)
        if not vals:
            return np.nan
        return max(vals)

    # ---- 节点航段预计算表（与载荷无关的几何量）----
    def build_leg_table(self):
        """节点 = O01 + 15 服务区。返回 dict[(from_key, to_key)] -> leg dict。
        key: 'O' 或 'S001'..'S015'。"""
        keys = ['O'] + sorted(self.areas.keys())
        table = {}
        op_alt = {}
        o = self.centers[self.OID]
        op_alt['O'] = o['alt']
        for sid in self.areas:
            op_alt[sid] = self.areas[sid]['alt'] + 30.0
        pts = {}
        pts['O'] = (o['lon'], o['lat'])
        for sid in self.areas:
            a = self.areas[sid]
            pts[sid] = (a['lon'], a['lat'])
        for a in keys:
            for b in keys:
                if a == b:
                    continue
                lon1, lat1 = pts[a]
                lon2, lat2 = pts[b]
                z1, z2 = op_alt[a], op_alt[b]
                d, cruise, h_up, h_down = segment_geometry(lon1, lat1, z1, lon2, lat2, z2, self)
                table[(a, b)] = {'d': d, 'cruise': cruise, 'h_up': h_up, 'h_down': h_down,
                                 'z1': z1, 'z2': z2}
        self.leg_table = table
        self.node_keys = keys
        self.op_alt = op_alt


# ----------------------------------------------------------------------------
# 航段时间与能耗
# ----------------------------------------------------------------------------
def Lg(g, q, data):
    """机型 g 携带载荷 q (kg) 的等效航程 (m)。"""
    p = data.uav_types[g]
    return p['L0'] - (p['L0'] - p['LF']) * (q / p['Q']) ** 1.5


def segment_geometry(lon1, lat1, z1, lon2, lat2, z2, data):
    """计算航段几何: 水平距离 d, 巡航海拔, 爬升/下降高度。"""
    d = data.dist_ll(lon1, lat1, lon2, lat2)
    cruise = data.corridor_max(lon1, lat1, lon2, lat2) + 50.0
    h_up = max(0.0, cruise - z1)
    h_down = max(0.0, cruise - z2)
    return d, cruise, h_up, h_down


def segment_time(d, h_up, h_down, g, data):
    p = data.uav_types[g]
    return h_up / p['v_up'] + d / p['v_c'] + h_down / p['v_down']


# ---- 基于预计算航段表的快速评估 ----
def leg_time_from_table(leg, g, data):
    p = data.uav_types[g]
    return leg['h_up'] / p['v_up'] + leg['d'] / p['v_c'] + leg['h_down'] / p['v_down']


def leg_energy_from_table(leg, q, g, data):
    p = data.uav_types[g]
    e_hor = p['E_use'] * leg['d'] / Lg(g, q, data)
    m_total = p['m0'] + q
    e_up = m_total * G * leg['h_up'] / (p['eta_up'] * J_PER_KWH)
    return e_hor + e_up


def route_profile_fast(data, keys, loads, g):
    """keys: ['O','S001',...,'O']；loads: 每航段载荷。返回 time/energy/legs。"""
    t_total = 0.0
    e_total = 0.0
    legs = []
    for i in range(len(keys) - 1):
        leg = data.leg_table[(keys[i], keys[i + 1])]
        q = loads[i]
        t_leg = leg_time_from_table(leg, g, data)
        e_leg = leg_energy_from_table(leg, q, g, data)
        t_total += t_leg
        e_total += e_leg
        legs.append({'from': keys[i], 'to': keys[i + 1], 'd': leg['d'],
                     'cruise': leg['cruise'], 'h_up': leg['h_up'], 'h_down': leg['h_down'],
                     't': t_leg, 'e': e_leg, 'q': q})
    return {'time': t_total, 'energy': e_total, 'legs': legs}


def segment_energy(d, h_up, q, g, data):
    """航段能耗 (kWh)：水平 + 爬升附加。"""
    p = data.uav_types[g]
    e_hor = p['E_use'] * d / Lg(g, q, data)
    m_total = p['m0'] + q
    e_up = m_total * G * h_up / (p['eta_up'] * J_PER_KWH)
    return e_hor + e_up


def flight_profile(nodes, loads, g, data):
    """
    计算一个运输架次的完整 profile。
    nodes: [(lon,lat,op_alt), ...]，第一个为 O01，随后为服务区（投送后载荷下降），最后回 O01。
    loads: 各航段 i->i+1 上的载荷 q (kg)，len = len(nodes)-1。
    返回: {'time': s, 'energy': kWh, 'legs': [...]}
    """
    p = data.uav_types[g]
    t_total = 0.0
    e_total = 0.0
    legs = []
    for i in range(len(nodes) - 1):
        lon1, lat1, z1 = nodes[i]
        lon2, lat2, z2 = nodes[i + 1]
        q = loads[i]
        d, cruise, h_up, h_down = segment_geometry(lon1, lat1, z1, lon2, lat2, z2, data)
        t_leg = segment_time(d, h_up, h_down, g, data)
        e_leg = segment_energy(d, h_up, q, g, data)
        t_total += t_leg
        e_total += e_leg
        legs.append({'from': i, 'to': i + 1, 'd': d, 'cruise': cruise,
                     'h_up': h_up, 'h_down': h_down, 't': t_leg, 'e': e_leg, 'q': q})
    return {'time': t_total, 'energy': e_total, 'legs': legs}


def flight_time_with_handover(profile, nbox, g, data):
    """架次总占用时间（不含排队）：准备+装载 + 飞行 + 交接。"""
    p = data.uav_types[g]
    return p['t_prep'] + nbox * p['t_load'] + profile['time'] + p['t_hand_base'] + nbox * p['t_hand_box']


# ----------------------------------------------------------------------------
# 电池充电
# ----------------------------------------------------------------------------
def charge_time(soc_end, T_full):
    """soc_end: 任务结束时 SOC（0~1 小数）。返回充电至 100% 所需时间 (s)。"""
    s = max(0.0, min(1.0, soc_end))
    if s < 0.90:
        return T_full * (0.65 * (0.90 - s) / 0.90 + 0.35)
    return T_full * 0.35 * (1.0 - s) / 0.10


# ----------------------------------------------------------------------------
# 通信链路
# ----------------------------------------------------------------------------
def _los_occluded(lon1, lat1, z1, lon2, lat2, z2, data, step=30.0):
    """向量化 LOS：返回是否被地形遮挡。"""
    d = data.dist_ll(lon1, lat1, lon2, lat2)
    n = max(2, int(math.ceil(d / step)))
    t = np.arange(1, n) / n
    lons = np.asarray(lon1 + (lon2 - lon1) * t)
    lats = np.asarray(lat1 + (lat2 - lat1) * t)
    z_line = z1 + (z2 - z1) * t
    z_terr = data.dem_at_many(lons, lats)
    bad = ~np.isnan(z_terr) & (z_terr > z_line + 1e-9)
    return bool(np.any(bad))


def fspl_db(f_mhz, d_km):
    return 32.45 + 20.0 * math.log10(f_mhz) + 20.0 * math.log10(d_km)


def path_loss(lon1, lat1, z1, lon2, lat2, z2, data):
    d2 = data.dist_ll(lon1, lat1, lon2, lat2)
    dz = (z2 - z1) if not (math.isnan(z1) or math.isnan(z2)) else 0.0
    d_km = math.sqrt(d2 * d2 + dz * dz) / 1000.0
    l = fspl_db(data.comm['f_MHz'], max(d_km, 1e-6))
    if _los_occluded(lon1, lat1, z1, lon2, lat2, z2, data):
        l += data.comm['Lobs']
    return l


def link_ok(lon1, lat1, z1, lon2, lat2, z2, lmax, data):
    """双向链路是否可用。"""
    if math.isnan(z1) or math.isnan(z2):
        return False
    return path_loss(lon1, lat1, z1, lon2, lat2, z2, data) <= lmax


def direct_ok(lon, lat, z, data):
    gw = data.gw_pt
    return link_ok(lon, lat, z, gw['lon'], gw['lat'], gw['z'], data.comm['Lmax_uav2gw'], data)


def relay_link_ok(lon, lat, z, rel_pt, data):
    """运输机-中继 接入链路。"""
    return link_ok(lon, lat, z, rel_pt[0], rel_pt[1], rel_pt[2],
                   data.comm['Lmax_uav2relay'], data)


def relay_backhaul_ok(rel_pt, data):
    gw = data.gw_pt
    return link_ok(rel_pt[0], rel_pt[1], rel_pt[2], gw['lon'], gw['lat'], gw['z'],
                   data.comm['Lmax_relay2gw'], data)


# ----------------------------------------------------------------------------
# 中继无人机时间与能耗
# ----------------------------------------------------------------------------
def relay_mission_time(relon, relat, rel_z, data):
    """中继架次：从 O01 到悬停点再返回。rel_z = 悬停海拔。
    返回 (t_out, t_back, t_total_flight)。不含准备/建链/服务/周转。"""
    o = data.centers[data.OID]
    z0 = o['alt']
    p = data.relay_type
    d, cruise, h_up, h_down = segment_geometry(o['lon'], o['lat'], z0, relon, relat, rel_z, data)
    t_out = h_up / p['v_up'] + d / p['v_c'] + h_down / p['v_down']
    d2, cruise2, h_up2, h_down2 = segment_geometry(relon, relat, rel_z, o['lon'], o['lat'], z0, data)
    t_back = h_up2 / p['v_up'] + d2 / p['v_c'] + h_down2 / p['v_down']
    return t_out, t_back, t_out + t_back


def relay_mission_energy(relon, relat, rel_z, data, t_serve):
    """中继架次能耗 (kWh)：
    水平巡航 = 巡航功率*巡航时间；爬升附加 = 总质量*g*h/eta；服务 = 悬停+通信附加功率*服务时间。
    """
    o = data.centers[data.OID]
    z0 = o['alt']
    p = data.relay_type
    d, cruise, h_up, h_down = segment_geometry(o['lon'], o['lat'], z0, relon, relat, rel_z, data)
    d2, cruise2, h_up2, h_down2 = segment_geometry(relon, relat, rel_z, o['lon'], o['lat'], z0, data)
    m_total = p['m_takeoff']
    e_cruise = p['P_cruise'] * (d + d2) / p['v_c'] / 3600.0
    e_up = m_total * G * (h_up + h_up2) / (p['eta_up'] * J_PER_KWH)
    e_serve = (p['P_hover'] + p['P_comm']) * t_serve / 3600.0
    return e_cruise + e_up + e_serve


def node_positions(nodes, loads, g, data, dt=10.0):
    """把运输架次展开为 (t, lon, lat, z, phase) 序列，用于通信判定。
    nodes: [(lon,lat,op_alt)]；loads: 各航段载荷。
    爬升在起点（水平不动），巡航水平移动，下降在终点。"""
    pts = [{'t': 0.0, 'lon': nodes[0][0], 'lat': nodes[0][1], 'z': nodes[0][2], 'phase': 'prep'}]
    t = 0.0
    p = data.uav_types[g]
    for i in range(len(nodes) - 1):
        lon1, lat1, z1 = nodes[i]
        lon2, lat2, z2 = nodes[i + 1]
        d, cruise, h_up, h_down = segment_geometry(lon1, lat1, z1, lon2, lat2, z2, data)
        # 爬升
        t_up = h_up / p['v_up']
        n = max(1, int(math.ceil(t_up / dt)))
        for k in range(1, n + 1):
            t += t_up / n
            pts.append({'t': t, 'lon': lon1, 'lat': lat1, 'z': z1 + h_up * k / n, 'phase': 'climb'})
        # 巡航
        t_cr = d / p['v_c']
        n = max(1, int(math.ceil(t_cr / dt)))
        for k in range(1, n + 1):
            t += t_cr / n
            pts.append({'t': t, 'lon': lon1 + (lon2 - lon1) * k / n,
                        'lat': lat1 + (lat2 - lat1) * k / n, 'z': cruise, 'phase': 'cruise'})
        # 下降
        t_dn = h_down / p['v_down']
        n = max(1, int(math.ceil(t_dn / dt)))
        for k in range(1, n + 1):
            t += t_dn / n
            pts.append({'t': t, 'lon': lon2, 'lat': lat2, 'z': cruise - h_down * k / n, 'phase': 'descent'})
    return pts


def main():
    data = Data()
    print('O01:', data.centers[data.OID])
    print('areas:', len(data.areas))
    print('boxes:', len(data.boxes))
    print('uav types:', {g: (p['Q'], p['E_use'], p['L0'], p['LF']) for g, p in data.uav_types.items()})
    print('uavs:', data.uavs)
    print('batteries:', data.batteries)
    print('relay:', data.relay_type)
    print('comm margins:', data.comm['Lmax_uav2gw'], data.comm['Lmax_uav2relay'], data.comm['Lmax_relay2gw'])
    # 测试一个往返 O01->S001->O01
    o = data.centers['O01']
    a = data.areas['S001']
    z1, z2 = o['alt'], a['alt'] + 30
    d, cruise, h_up, h_down = segment_geometry(o['lon'], o['lat'], z1, a['lon'], a['lat'], z2, data)
    print('S001 leg: d=%.0fm cruise=%.1f h_up=%.1f h_down=%.1f' % (d, cruise, h_up, h_down))
    prof = flight_profile([(o['lon'], o['lat'], z1), (a['lon'], a['lat'], z2), (o['lon'], o['lat'], z1)],
                          [10.0, 0.0], 'A', data)
    print('round trip A q=10kg: t=%.0fs e=%.3fkWh' % (prof['time'], prof['energy']))
    print('charge time 0.5 ->', charge_time(0.5, 1800))


if __name__ == '__main__':
    main()
