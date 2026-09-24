# -*- coding: utf-8 -*-
"""把 Q1-Q4 结果写入 结果提交模板.xlsx 的六个工作表。"""
import json
import os
import sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from core import Data

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(BASE)
OUT = os.path.join(BASE, 'results')
TPL = os.path.join(ROOT, '第二十三届中国研究生数学建模竞赛 - 中文题目', '中文题目', 'D题', '结果提交模板.xlsx')
DST = os.path.join(ROOT, '结果提交模板_填写.xlsx')

data = Data()
POS = {'W': (109.2103, 23.047134, 676.5),
       'E': (109.276017, 23.019401, 542.3),
       'N': (109.238171, 23.077841, 496.1)}
UAV_IDS = {'A': ['U01', 'U02', 'U03', 'U04'], 'B': ['U05', 'U06'], 'C': ['U07', 'U08']}
BAT_IDS = {'A': ['A-01', 'A-02', 'A-03', 'A-04', 'A-05', 'A-06'],
           'B': ['B-01', 'B-02', 'B-03', 'B-04'],
           'C': ['C-01', 'C-02', 'C-03', 'C-04']}
MODEL_OF = dict(data.uavs)

head_fill = PatternFill('solid', fgColor='DDEBF7')
head_font = Font(bold=True)


def set_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal='center', vertical='center')


def fill_q1(ws):
    r = json.load(open(os.path.join(OUT, 'p1_results.json'), encoding='utf-8'))
    rows = []
    k = 1
    for area, g in r['grouping'].items():
        for d in g['detail']:
            boxes = d['boxes']
            mass = sum(data.boxes[b]['mass'] for b in boxes)
            vol = sum(data.boxes[b]['vol'] for b in boxes)
            rows.append(['P1-%02d' % k, area, d['model'], '、'.join(boxes), round(mass, 1),
                         round(vol, 3), round(d['time'], 1), round(d['energy'], 3),
                         round((1 - d['energy'] / data.uav_types[d['model']]['E_use']) * 100, 1)])
            k += 1
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    print('Q1 rows:', len(rows))


def fill_q2(ws_flight, ws_box):
    # Q2 工作表填写问题二均衡方案（24 架次），来源为 p2_pareto.json 的 balanced 键
    r = json.load(open(os.path.join(OUT, 'p2_pareto.json'), encoding='utf-8'))['balanced']
    rows_f = []
    rows_b = []
    for f in sorted(r['flights'], key=lambda x: x['fid']):
        seq = '→'.join(s for s, _ in f['route'])
        rows_f.append([f['fid'], f['uav'], f['model'], f['battery'],
                       round(f['start'], 1), seq, round(f['return'], 1), round(f['energy'], 3)])
    for d in r['deliveries']:
        rows_b.append([d['box'], d['fid'], d['area'], round(d['t'], 1)])
    for i, row in enumerate(rows_f, start=2):
        for j, v in enumerate(row, start=1):
            ws_flight.cell(row=i, column=j, value=v)
    for i, row in enumerate(rows_b, start=2):
        for j, v in enumerate(row, start=1):
            ws_box.cell(row=i, column=j, value=v)
    print('Q2 flights:', len(rows_f), ' boxes:', len(rows_b))


def fill_q3(ws_sortie, ws_comm):
    r = json.load(open(os.path.join(OUT, 'p3_final.json'), encoding='utf-8'))
    sched = r['schedule']
    from core import relay_mission_time
    # 中继架次表
    rows = []
    r1_n = r2_n = 0
    for k, sg in enumerate(r['sorties'], start=1):
        relay = sg['relay']          # R01 / R02
        pos = sg['pos']
        lo, la, z = POS[pos]
        t_out, t_back, _ = relay_mission_time(lo, la, z, data)
        t0, t1 = sg['t0'], sg['t1']
        dep = t0 - t_out - 30.0
        if relay == 'R01':
            r1_n += 1
            rid = 'R1-%02d' % r1_n
            uav = '中继1'
        else:
            r2_n += 1
            rid = 'R2-%02d' % r2_n
            uav = '中继2'
        rows.append([rid, uav, '能源-%02d' % k,
                     round(dep, 1), lo, la, z,
                     round(t0, 1), round(t1, 1), round(t1 + t_back, 1), round(sg['energy'], 3)])
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws_sortie.cell(row=i, column=j, value=v)
    # 通信保障表：逐架次逐服务区阶段
    rows_c = []
    for f in sorted(r['flights'], key=lambda x: x['fid']):
        for sid, bid in f['route']:
            pass
    # 直接用 p3_final 的 missions（需要重建：从架次轨迹判定）
    from p2_solve import Flight
    from p3_gaps import flight_trajectory
    from core import direct_ok
    flights = [Flight(fj['fid'], [(s, list(b)) for s, b in fj['route']], fj['model'], data)
               for fj in r['flights']]
    sortie_of = {}
    for sg in r['sorties']:
        for mf in sg['missions']:
            sortie_of[mf] = sg
    for f in flights:
        start = sched[str(f.fid)]['start']
        pts, _ = flight_trajectory(f, start)
        # 逐点归属最近服务区（用于阶段划分）
        for pt in pts:
            pt['area'] = min(data.areas,
                             key=lambda s: data.dist_ll(pt['lon'], pt['lat'],
                                                        data.areas[s]['lon'], data.areas[s]['lat']))
        # 按航线顺序输出每个服务区阶段（含转运直连段）
        visited = {}
        order = []
        for pt in pts:
            a = pt.get('area')
            if a and a not in visited:
                visited[a] = True
                order.append(a)
        for a in order:
            sp = [pt for pt in pts if pt.get('area') == a and pt['phase'] != 'prep']
            if not sp:
                continue
            t0 = min(pt['t'] for pt in sp)
            t1 = max(pt['t'] for pt in sp)
            need = [pt for pt in sp if not direct_ok(pt['lon'], pt['lat'], pt['z'], data)]
            if not need:
                rows_c.append([f.fid, a, round(t0, 1), round(t1, 1), '直连', '—'])
            else:
                sg = None
                for s2 in r['sorties']:
                    if min(t0, min(pt['t'] for pt in need)) >= s2['t0'] - 1 and \
                       max(t1, max(pt['t'] for pt in need)) <= s2['t1'] + 1:
                        sg = s2
                rid = '—'
                if sg:
                    rid = ('R1-%02d' if sg['relay'] == 'R01' else 'R2-%02d') % (
                        r['sorties'].index(sg) + 1)
                rows_c.append([f.fid, a, round(t0, 1), round(t1, 1), '中继', rid])
    for i, row in enumerate(rows_c, start=2):
        for j, v in enumerate(row, start=1):
            ws_comm.cell(row=i, column=j, value=v)
    print('Q3 sorties:', len(rows), ' comm rows:', len(rows_c))


def fill_q4(ws):
    r = json.load(open(os.path.join(OUT, 'p4_results.json'), encoding='utf-8'))
    rows = []
    for K in ('3', '2'):
        blk = r[K]
        for g in blk['groups']:
            au = g['alloc_uav']; ab = g['alloc_bat']
            rows.append([int(K), g['name'], '、'.join(g['areas']),
                         au.get('A', 0), au.get('B', 0), au.get('C', 0),
                         ab.get('A', 0), ab.get('B', 0), ab.get('C', 0),
                         g['alloc_relay'], g['alloc_comp']])
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    print('Q4 rows:', len(rows))


def main():
    wb = openpyxl.load_workbook(TPL)
    fill_q1(wb['Q1_单点组批'])
    fill_q2(wb['Q2_运输架次'], wb['Q2_逐箱交付'])
    fill_q3(wb['Q3_中继架次'], wb['Q3_通信保障'])
    fill_q4(wb['Q4_分区配置'])
    for ws in wb.worksheets:
        set_header(ws, ws.max_column)
    wb.save(DST)
    print('saved', DST)


if __name__ == '__main__':
    main()
